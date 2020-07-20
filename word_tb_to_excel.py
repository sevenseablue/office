from docx import Document
import xlsxwriter
import sys

path = sys.argv[1]
if path.endswith(".docx"):
    suff_ind = -5
elif path.endswith(".doc"):
    suff_ind = -4
else:
    suff_ind = -4
path_xls = path[:suff_ind] + ".xlsx"
print("准备将word中的表格写入文件" + path_xls + "中")
workbook = xlsxwriter.Workbook(path_xls)

doc = Document(path)
for table in doc.tables:
    worksheet1 = workbook.add_worksheet()

    for r_ind, row in enumerate(table.rows):
        row_text = [c.text for c in row.cells]
        print(row_text)
        # TODO: print row in Excel with xlsxwriter or openpyxl

        # For xlsxwriter: use worksheet.write_row()
        # See https://xlsxwriter.readthedocs.io/worksheet.html for details
        for c_ind, cell in enumerate(row_text):
            worksheet1.write(r_ind, c_ind, cell)


        # For openpyxl: see doc in https://openpyxl.readthedocs.io/en/stable/tutorial.htm

workbook.close()
print("写入完毕")

